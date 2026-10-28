#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
使照勘驗缺失照片自動匯入 Excel 工具
======================================

用途：
    讀取「使照勘驗缺失紀錄表.xlsx」的「勘驗紀錄」分頁，依編號（D-001、D-002...）
    到「照片」資料夾找出對應的改善前／改善後照片，壓縮後嵌入 Excel，
    每個缺失（或每組缺失）產生一個分頁，並在主表加上可點擊的照片連結。

用法：
    python build_photo_report.py [來源Excel路徑] [照片資料夾路徑] [--output 輸出路徑]

    不帶參數時，預設使用目前資料夾下的：
        使照勘驗缺失紀錄表.xlsx
        照片/

詳細說明請見 README.md。
"""

from __future__ import annotations

import argparse
import io
import math
import re
import sys
import traceback
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    print("❌ 錯誤：找不到 openpyxl 套件，請先執行：pip install openpyxl")
    sys.exit(1)

try:
    from PIL import Image as PILImage
    from PIL import ImageOps
except ImportError:
    print("❌ 錯誤：找不到 Pillow 套件，請先執行：pip install Pillow")
    sys.exit(1)

# HEIC/HEIF 支援為選用套件；沒安裝時不靜默跳過，而是在真的遇到 HEIC 檔案時明確報錯
HEIF_AVAILABLE = False
try:
    import pillow_heif

    pillow_heif.register_heif_opener()
    HEIF_AVAILABLE = True
except ImportError:
    HEIF_AVAILABLE = False


# ======================================================================
# 可調整常數
# ======================================================================

# 分頁數量門檻：有照片的缺失項目數量超過此值時，改為依「查驗項目」分組建立分頁，
# 而不是一項一頁（避免 Excel 分頁過多卡死）。可依需要調整。
MAX_INDIVIDUAL_SHEETS = 60

# 超過門檻時的分組依據：'item' = 依查驗項目（D欄，20 種，分組數可控）
#                        'location' = 依位置/樓層（C欄，文字較自由，分組數不固定）
GROUP_BY = "item"
GROUP_BY_LABEL = {"item": "查驗項目", "location": "位置/樓層"}[GROUP_BY]

# 照片壓縮參數
MAX_DIMENSION = 1600  # 長邊縮放上限（像素），小圖不放大
JPEG_QUALITY = 80  # 預設 JPEG 品質
MIN_JPEG_QUALITY = 40  # 品質下限（避免越壓越糊也救不回容量時無限迴圈）
TARGET_MAX_BYTES = 300 * 1024  # 單張目標上限 300KB（超過會逐步降低品質重試）

# 支援的照片副檔名
SUPPORTED_EXT = {".jpg", ".jpeg", ".png", ".heic", ".heif"}

# 分頁版面：顯示用常數
PHOTO_COL_WIDTH = 22  # 照片分頁 A~D 欄寬（字元數）
PHOTO_DISPLAY_WIDTH = 300  # 相片在 Excel 中顯示的寬度（像素，非原始解析度）
PX_TO_PT = 0.75  # 96dpi 下 1pt = 1.333px，換算列高用
DEFAULT_ROW_HEIGHT_PT = 15  # Excel 預設列高（pt）

# 配色（沿用主表風格）
TITLE_BLUE = "1F3864"
HEADER_BLUE = "2E75B6"
LIGHT_BLUE = "D6E4F0"
HEADER_GREY = "595959"
MUTED_GREY = "A6A6A6"

# 狀態配色（與主表 K 欄條件式格式一致）：(底色, 字色)
STATUS_COLORS = {
    "已改善": ("C6EFCE", "006100"),
    "待複查": ("DDEBF7", "1F4E79"),
    "改善中": ("FFEB9C", "9C5700"),
    "逾期": ("FFC7CE", "9C0006"),
}

MAIN_SHEET_NAME = "勘驗紀錄"
PROTECTED_SHEETS = {MAIN_SHEET_NAME, "統計", "查驗項目參考", "選單"}

# 本工具自動產生的分頁命名規則：D-001...（逐項）或 G-xxx（分組）。
# 每次重新執行前會先清除舊的產生分頁，避免重複增生 D-001 (2) 這類分頁。
GENERATED_SHEET_PATTERN = re.compile(r"^(D-\d+|G-.+)$")

DEBUG = False


# ======================================================================
# 共用小工具
# ======================================================================

def die(msg: str) -> None:
    """印出明確錯誤訊息並結束程式（不印出醜醜的 traceback）。"""
    print(f"\n❌ 錯誤：{msg}\n", file=sys.stderr)
    sys.exit(1)


def font(size=10, bold=False, color="000000", underline=None, name="Arial"):
    return Font(name=name, size=size, bold=bold, color=color, underline=underline)


def solid_fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


_THIN_SIDE = Side(style="thin", color="BFBFBF")
BORDER_THIN = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 不可出現在 Excel 分頁名稱中的字元（含單引號，避免破壞 HYPERLINK 公式的引號）
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:\*\?/\\']")


def sanitize_sheet_name(name: str, max_len: int = 31) -> str:
    """把文字轉成合法的 Excel 分頁名稱：移除禁用字元、限制長度。"""
    name = _INVALID_SHEET_CHARS.sub("_", str(name)).strip()
    if not name:
        name = "Sheet"
    return name[:max_len]


def unique_sheet_name(base: str, used_names: set) -> str:
    """確保分頁名稱不重複（同一次執行內）。"""
    name = sanitize_sheet_name(base)
    if name not in used_names:
        used_names.add(name)
        return name
    i = 2
    while True:
        suffix = f"_{i}"
        candidate = sanitize_sheet_name(base)[: 31 - len(suffix)] + suffix
        if candidate not in used_names:
            used_names.add(candidate)
            return candidate
        i += 1


# ======================================================================
# 第一步：讀取主表「勘驗紀錄」中的缺失資料
# ======================================================================

def compute_status(review_result, deadline, completed, today: date) -> str:
    """
    依照主表 K 欄的邏輯在 Python 端重新計算狀態，不依賴 Excel 快取的計算結果
    （因為來源檔案若還沒在 Excel 中開啟過，openpyxl 讀不到公式計算後的值）。

    對應主表公式：
        =IF($E6="","",IF($J6="合格","已改善",
            IF(AND($H6<>"",$H6<TODAY(),$I6=""),"逾期",
                IF($I6<>"","待複查","改善中"))))
    """
    if review_result == "合格":
        return "已改善"
    deadline_is_overdue = isinstance(deadline, datetime) and deadline.date() < today
    if deadline_is_overdue and not completed:
        return "逾期"
    if completed:
        return "待複查"
    return "改善中"


def read_defects(ws) -> list[dict]:
    """讀取「勘驗紀錄」分頁中所有 E 欄（缺失記錄）有值的資料列。"""
    defects = []
    last_row = max(ws.max_row, 125)
    today = date.today()
    for row in range(6, last_row + 1):
        e_val = ws.cell(row=row, column=5).value
        if e_val is None or str(e_val).strip() == "":
            continue
        responsible = ws.cell(row=row, column=7).value
        deadline = ws.cell(row=row, column=8).value
        completed = ws.cell(row=row, column=9).value
        review_result = ws.cell(row=row, column=10).value
        defects.append(
            {
                "row": row,
                "id": f"D-{row - 5:03d}",  # 對應主表 A 欄公式：D- + (ROW()-5) 補零 3 位
                "location": ws.cell(row=row, column=3).value or "",
                "item": ws.cell(row=row, column=4).value or "",
                "defect_text": str(e_val),
                "responsible": responsible or "",
                "deadline": deadline,
                "status": compute_status(review_result, deadline, completed, today),
            }
        )
    return defects


def clear_previous_generated_sheets(wb) -> None:
    """
    重跑前先清掉上次工具自動產生的分頁（D-xxx / G-xxx），
    確保「重複執行不會重複增生分頁」，且一律以最新資料重新產生。
    4 個原始分頁（勘驗紀錄/統計/查驗項目參考/選單）絕不會被清除。
    """
    for name in list(wb.sheetnames):
        if name in PROTECTED_SHEETS:
            continue
        if GENERATED_SHEET_PATTERN.match(name):
            del wb[name]


# ======================================================================
# 第二步：掃描照片資料夾
# ======================================================================

def discover_photos(photo_root: Path, defect_id: str):
    """
    依規則掃描 照片/<編號>/ 底下的照片。

    回傳值：
        None              -> 找不到該編號的資料夾
        []                -> 資料夾存在，但沒有任何支援格式的照片
        [(分類, Path), ...] -> 找到的照片，分類為「改善前」或「改善後」
    """
    folder = photo_root / defect_id
    if not folder.is_dir():
        return None

    items = []
    sub_before = folder / "改善前"
    sub_after = folder / "改善後"
    has_subfolder = sub_before.is_dir() or sub_after.is_dir()

    if has_subfolder:
        if sub_before.is_dir():
            for p in sorted(sub_before.iterdir()):
                if p.is_file() and p.suffix.lower() in SUPPORTED_EXT:
                    items.append(("改善前", p))
        if sub_after.is_dir():
            for p in sorted(sub_after.iterdir()):
                if p.is_file() and p.suffix.lower() in SUPPORTED_EXT:
                    items.append(("改善後", p))
    else:
        # 沒有子資料夾：資料夾內所有照片一律視為「改善前」
        for p in sorted(folder.iterdir()):
            if p.is_file() and p.suffix.lower() in SUPPORTED_EXT:
                items.append(("改善前", p))

    return items


# ======================================================================
# 第三步：照片壓縮
# ======================================================================

def _extract_exif_datetime(img: "PILImage.Image"):
    """取出照片的拍攝時間（EXIF DateTimeOriginal，沒有則退回 DateTime）。"""
    try:
        exif = img.getexif()
        if not exif:
            return None
        dt_str = exif.get(306)  # DateTime
        try:
            exif_ifd = exif.get_ifd(0x8769)  # Exif IFD
            dt_str = exif_ifd.get(36867, dt_str)  # DateTimeOriginal 優先
        except Exception:
            pass
        if dt_str:
            return datetime.strptime(str(dt_str), "%Y:%m:%d %H:%M:%S")
    except Exception:
        pass
    return None


def compress_image(path: Path) -> dict:
    """
    讀取單張照片、校正方向、縮圖、壓縮為 JPEG，回傳可直接嵌入 Excel 的資料。

    重點（對應需求）：
      1. 一律先套用 EXIF Orientation 再做任何縮放/儲存，避免直式照片躺平。
      2. 長邊縮到 MAX_DIMENSION，原圖較小則不放大。
      3. JPEG 品質從 JPEG_QUALITY 開始，若超過 TARGET_MAX_BYTES 則逐步降低品質重試。
      4. 全程只在記憶體（BytesIO）中處理，不寫任何暫存檔到磁碟，自然不需要清暫存檔。
    """
    ext = path.suffix.lower()
    if ext in (".heic", ".heif") and not HEIF_AVAILABLE:
        raise RuntimeError(
            "這是 HEIC/HEIF 照片，但目前環境沒有安裝 pillow-heif，無法轉檔。"
            "請先執行：pip install pillow-heif"
        )

    orig_size = path.stat().st_size

    try:
        img = PILImage.open(path)
        img.load()
    except Exception as e:  # 照片損毀、格式錯誤等
        raise RuntimeError(f"照片無法開啟或已損毀（{e}）") from e

    exif_dt = _extract_exif_datetime(img)

    # 先校正方向，才做任何後續處理
    img = ImageOps.exif_transpose(img)

    # 統一轉成 RGB（去除透明度，避免存成 JPEG 時出錯）
    if img.mode in ("RGBA", "LA"):
        background = PILImage.new("RGB", img.size, (255, 255, 255))
        background.paste(img.convert("RGBA"), mask=img.convert("RGBA").split()[-1])
        img = background
    elif img.mode == "P":
        img = img.convert("RGBA")
        background = PILImage.new("RGB", img.size, (255, 255, 255))
        background.paste(img, mask=img.split()[-1])
        img = background
    elif img.mode != "RGB":
        img = img.convert("RGB")

    w, h = img.size
    long_side = max(w, h)
    if long_side > MAX_DIMENSION:
        scale = MAX_DIMENSION / long_side
        img = img.resize((max(1, round(w * scale)), max(1, round(h * scale))), PILImage.LANCZOS)

    buf = io.BytesIO()
    quality = JPEG_QUALITY
    while True:
        buf.seek(0)
        buf.truncate(0)
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        size = buf.tell()
        if size <= TARGET_MAX_BYTES or quality <= MIN_JPEG_QUALITY:
            break
        quality -= 5
    buf.seek(0)

    # 重新用 BytesIO 開啟一次，讓 openpyxl 可以直接用這個物件的 fp 讀取影像資料。
    # 注意：不要呼叫 thumb.load()，否則 Pillow 會把 fp 設為 None，openpyxl 讀不到資料。
    thumb = PILImage.open(buf)

    return {
        "thumb": thumb,
        "buf": buf,  # 保留參照避免被 GC 回收，直到 wb.save() 完成
        "width": thumb.width,
        "height": thumb.height,
        "orig_size": orig_size,
        "new_size": size,
        "exif_dt": exif_dt,
    }


# ======================================================================
# 第四步：建立缺失照片分頁
# ======================================================================

def setup_photo_sheet_columns(ws) -> None:
    for col in ("A", "B", "C", "D"):
        ws.column_dimensions[col].width = PHOTO_COL_WIDTH


def _write_info_row(ws, r: int, label: str, value, wrap_tall: bool = False, date_fmt: bool = False):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    label_cell = ws.cell(row=r, column=1, value=label)
    label_cell.font = font(bold=True)
    label_cell.fill = solid_fill(LIGHT_BLUE)
    label_cell.alignment = ALIGN_CENTER
    label_cell.border = BORDER_THIN

    value_cell = ws.cell(row=r, column=2, value=value)
    value_cell.border = BORDER_THIN
    value_cell.alignment = ALIGN_LEFT_WRAP
    if date_fmt and isinstance(value, datetime):
        value_cell.number_format = "yyyy/mm/dd"
    for c in (3, 4):
        ws.cell(row=r, column=c).border = BORDER_THIN

    ws.row_dimensions[r].height = 46 if wrap_tall else 20


def write_defect_block(ws, start_row: int, defect: dict, photos: list[dict]) -> int:
    """
    在工作表 start_row 開始寫入一個缺失的完整版面（標題列 + 資料區 + 照片區）。
    回傳最後使用到的列號（end_row），呼叫端可接著往下排下一個缺失區塊。
    """
    r = start_row

    # --- 標題列：缺失編號 + 右上角回主表連結 ---
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    title_cell = ws.cell(row=r, column=1, value=f"缺失照片：{defect['id']}")
    title_cell.font = font(size=13, bold=True, color="FFFFFF")
    title_cell.fill = solid_fill(TITLE_BLUE)
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in (2, 3):
        ws.cell(row=r, column=c).fill = solid_fill(TITLE_BLUE)

    back_cell = ws.cell(
        row=r,
        column=4,
        value=f"=HYPERLINK(\"#'{MAIN_SHEET_NAME}'!A{defect['row']}\",\"← 回勘驗紀錄\")",
    )
    back_cell.font = font(size=9, color="FFFFFF", underline="single")
    back_cell.fill = solid_fill(TITLE_BLUE)
    back_cell.alignment = ALIGN_CENTER
    ws.row_dimensions[r].height = 26
    r += 1

    # --- 基本資料區（從主表帶入） ---
    _write_info_row(ws, r, "編號", defect["id"]); r += 1
    _write_info_row(ws, r, "位置（樓層/區域）", defect["location"]); r += 1
    _write_info_row(ws, r, "查驗項目", defect["item"]); r += 1
    _write_info_row(ws, r, "缺失記錄", defect["defect_text"], wrap_tall=True); r += 1
    _write_info_row(ws, r, "責任單位/廠商", defect["responsible"]); r += 1
    _write_info_row(ws, r, "改善期限", defect["deadline"], date_fmt=True); r += 1

    # --- 狀態列（配色呼應主表圖例） ---
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    label_cell = ws.cell(row=r, column=1, value="狀態")
    label_cell.font = font(bold=True)
    label_cell.fill = solid_fill(LIGHT_BLUE)
    label_cell.alignment = ALIGN_CENTER
    label_cell.border = BORDER_THIN

    status_fill, status_font_color = STATUS_COLORS.get(defect["status"], ("FFFFFF", "000000"))
    value_cell = ws.cell(row=r, column=2, value=defect["status"])
    value_cell.font = font(bold=True, color=status_font_color)
    value_cell.fill = solid_fill(status_fill)
    value_cell.alignment = ALIGN_CENTER
    value_cell.border = BORDER_THIN
    for c in (3, 4):
        cc = ws.cell(row=r, column=c)
        cc.fill = solid_fill(status_fill)
        cc.border = BORDER_THIN
    ws.row_dimensions[r].height = 20
    r += 1

    r += 1  # 間距列

    # --- 照片區標題 ---
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    photo_header = ws.cell(row=r, column=1, value=f"現場照片（共 {len(photos)} 張）")
    photo_header.font = font(bold=True, color="FFFFFF")
    photo_header.fill = solid_fill(HEADER_BLUE)
    photo_header.alignment = ALIGN_CENTER
    ws.row_dimensions[r].height = 22
    r += 1

    end_row = build_photo_grid(ws, r, photos)
    return end_row


def build_photo_grid(ws, start_row: int, photos: list[dict]) -> int:
    """以 2 欄網格排列照片，每張照片下方標註分類／檔名／拍攝時間，避免疊字或超框。"""
    r = start_row
    col_pairs = [("A", "B"), ("C", "D")]

    for i in range(0, len(photos), 2):
        pair = photos[i : i + 2]
        disp_sizes = []
        for p in pair:
            disp_w = PHOTO_DISPLAY_WIDTH
            disp_h = round(disp_w * p["height"] / p["width"]) if p["width"] else disp_w
            disp_sizes.append((disp_w, disp_h))

        max_disp_h = max(h for _, h in disp_sizes)
        rows_needed = max(1, math.ceil(max_disp_h * PX_TO_PT / DEFAULT_ROW_HEIGHT_PT))
        box_top, box_bottom = r, r + rows_needed - 1

        for j, photo in enumerate(pair):
            c1, c2 = col_pairs[j]
            ws.merge_cells(f"{c1}{box_top}:{c2}{box_bottom}")
            for rr in range(box_top, box_bottom + 1):
                for cc in (c1, c2):
                    ws[f"{cc}{rr}"].border = BORDER_THIN
            ws[f"{c1}{box_top}"].fill = solid_fill("FFFFFF")

            disp_w, disp_h = disp_sizes[j]
            xl_img = XLImage(photo["thumb"])
            xl_img.width = disp_w
            xl_img.height = disp_h
            ws.add_image(xl_img, f"{c1}{box_top}")

            for rr in range(box_top, box_bottom + 1):
                if (ws.row_dimensions[rr].height or 0) < DEFAULT_ROW_HEIGHT_PT:
                    ws.row_dimensions[rr].height = DEFAULT_ROW_HEIGHT_PT

        caption_row = box_bottom + 1
        for j, photo in enumerate(pair):
            c1, c2 = col_pairs[j]
            ws.merge_cells(f"{c1}{caption_row}:{c2}{caption_row}")
            cap_cell = ws[f"{c1}{caption_row}"]
            dt_str = photo["exif_dt"].strftime("%Y/%m/%d %H:%M") if photo["exif_dt"] else "—"
            cap_cell.value = f"{photo['category']}｜{photo['filename']}｜{dt_str}"
            cap_cell.font = font(size=9, color="595959")
            cap_cell.alignment = ALIGN_CENTER
        ws.row_dimensions[caption_row].height = 16

        r = caption_row + 2  # 預留一列間距給下一排照片

    return r - 2 if photos else r - 1


# ======================================================================
# 第五步：主表加上照片連結
# ======================================================================

def update_main_sheet_links(ws, defects: list[dict], link_map: dict) -> None:
    header = ws.cell(row=5, column=13, value="照片")
    header.font = font(bold=True, color="FFFFFF")
    header.fill = solid_fill(HEADER_GREY)
    header.alignment = ALIGN_CENTER
    header.border = BORDER_THIN
    ws.column_dimensions["M"].width = 14

    for d in defects:
        row = d["row"]
        cell = ws.cell(row=row, column=13)
        cell.border = BORDER_THIN
        cell.alignment = ALIGN_CENTER
        info = link_map.get(d["id"])
        if info:
            sheet_name, anchor, count = info
            cell.value = f"=HYPERLINK(\"#'{sheet_name}'!{anchor}\",\"📷 {count}張\")"
            cell.font = font(color="0563C1", underline="single")
        else:
            cell.value = "—"
            cell.font = font(color=MUTED_GREY)

    last_row = max(ws.max_row, 125)
    ws.auto_filter.ref = f"A5:M{last_row}"


# ======================================================================
# 主流程
# ======================================================================

def resolve_output_path(excel_path: str, output_arg: str | None) -> Path:
    if output_arg:
        return Path(output_arg)
    src = Path(excel_path)
    date_str = date.today().strftime("%Y%m%d")
    return src.with_name(f"{src.stem}_含照片_{date_str}{src.suffix}")


def parse_args():
    p = argparse.ArgumentParser(description="使照勘驗缺失照片自動匯入 Excel 工具")
    p.add_argument(
        "excel",
        nargs="?",
        default="使照勘驗缺失紀錄表.xlsx",
        help="來源 Excel 檔案路徑（預設：目前資料夾下的 使照勘驗缺失紀錄表.xlsx）",
    )
    p.add_argument(
        "photos",
        nargs="?",
        default="照片",
        help="照片資料夾路徑（預設：目前資料夾下的 照片/）",
    )
    p.add_argument("--output", "-o", default=None, help="輸出檔案路徑（預設自動依日期命名）")
    p.add_argument("--debug", action="store_true", help="顯示完整錯誤堆疊，方便除錯")
    return p.parse_args()


def print_summary(
    excel_path, photos_root, out_path, defects, defect_photo_map, photo_count,
    total_orig, total_new, sheet_mode, group_count, warnings, errors,
):
    line = "=" * 56
    print(line)
    print("使照勘驗缺失照片匯入工具 - 執行摘要")
    print(line)
    print(f"來源檔案：{excel_path}")
    print(f"照片資料夾：{photos_root}")
    print("-" * 56)
    print(f"缺失總數（E欄有值）：{len(defects)} 筆")
    print(f"找到照片並建立分頁的缺失：{len(defect_photo_map)} 筆")
    print(f"有效照片總數：{photo_count} 張")
    if sheet_mode == "individual":
        print("分頁建立方式：逐項建立（每個缺失一個分頁）")
    elif sheet_mode == "group":
        print(f"分頁建立方式：依「{GROUP_BY_LABEL}」分組建立，共 {group_count} 個分頁")
    else:
        print("分頁建立方式：（無，沒有任何缺失找到照片）")
    print("-" * 56)
    if photo_count:
        print(
            f"照片壓縮：原始總容量 {total_orig / 1024 / 1024:.2f} MB → "
            f"壓縮後 {total_new / 1024 / 1024:.2f} MB "
            f"（節省 {(1 - total_new / total_orig) * 100:.1f}%）"
        )
    print("-" * 56)
    print(f"⚠ 警告（{len(warnings)} 項）：")
    for w in warnings:
        print(f"  - {w}")
    print(f"❌ 錯誤（{len(errors)} 項）：")
    for e in errors:
        print(f"  - {e}")
    print("-" * 56)
    print(f"輸出檔案：{out_path}")
    print(line)


def main():
    global DEBUG
    args = parse_args()
    DEBUG = args.debug

    excel_path = Path(args.excel)
    photos_root = Path(args.photos)

    if not excel_path.exists():
        die(f"找不到來源 Excel 檔案：{excel_path}")

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=False)
    except Exception as e:
        die(f"無法開啟 Excel 檔案：{excel_path}\n原因：{e}")

    if MAIN_SHEET_NAME not in wb.sheetnames:
        die(f"來源檔案缺少『{MAIN_SHEET_NAME}』分頁，請確認檔案格式正確、未被改名。")

    main_ws = wb[MAIN_SHEET_NAME]

    # 重跑保護：先清掉舊的自動產生分頁，確保不會重複增生
    clear_previous_generated_sheets(wb)

    defects = read_defects(main_ws)
    warnings: list[str] = []
    errors: list[str] = []

    if not defects:
        warnings.append(
            f"主表「{MAIN_SHEET_NAME}」中沒有任何缺失記錄（E欄皆為空白），本次沒有項目可處理。"
        )

    if not photos_root.exists():
        die(f"找不到照片資料夾：{photos_root}")

    valid_ids = {d["id"] for d in defects}
    for child in sorted(photos_root.iterdir()):
        if child.is_dir() and child.name not in valid_ids:
            warnings.append(f"資料夾「{child.name}」在主表查無對應編號，已忽略。")

    # --- 掃描照片 + 壓縮 ---
    defect_photo_map: "OrderedDict[str, list[dict]]" = OrderedDict()
    total_orig = total_new = photo_count = 0

    for d in defects:
        raw = discover_photos(photos_root, d["id"])
        if raw is None:
            warnings.append(f"{d['id']}：找不到對應照片資料夾，未建立分頁。")
            continue
        if not raw:
            warnings.append(f"{d['id']}：照片資料夾存在，但沒有找到支援格式的照片，未建立分頁。")
            continue

        processed = []
        for category, path in raw:
            try:
                info = compress_image(path)
            except Exception as e:
                errors.append(f"{d['id']}／{path.name}：{e}")
                continue
            info["category"] = category
            info["filename"] = path.name
            processed.append(info)
            total_orig += info["orig_size"]
            total_new += info["new_size"]
            photo_count += 1

        if not processed:
            warnings.append(f"{d['id']}：資料夾內的照片皆無法讀取，未建立分頁。")
            continue

        defect_photo_map[d["id"]] = processed

    defects_with_photos = [d for d in defects if d["id"] in defect_photo_map]

    # --- 建立分頁 ---
    link_map: dict[str, tuple[str, str, int]] = {}
    used_names = set(wb.sheetnames)
    sheet_mode = "none"
    group_count = 0

    if defects_with_photos and len(defects_with_photos) <= MAX_INDIVIDUAL_SHEETS:
        sheet_mode = "individual"
        for d in defects_with_photos:
            photos = defect_photo_map[d["id"]]
            sheet_name = unique_sheet_name(d["id"], used_names)
            ws = wb.create_sheet(sheet_name)
            setup_photo_sheet_columns(ws)
            write_defect_block(ws, 1, d, photos)
            link_map[d["id"]] = (sheet_name, "A1", len(photos))

    elif defects_with_photos:
        sheet_mode = "group"
        print(
            f"⚠ 有照片的缺失項目共 {len(defects_with_photos)} 筆，超過門檻 "
            f"{MAX_INDIVIDUAL_SHEETS} 筆，改為依「{GROUP_BY_LABEL}」分組建立分頁。"
        )
        groups: "OrderedDict[str, list[dict]]" = OrderedDict()
        for d in defects_with_photos:
            key = (d["item"] if GROUP_BY == "item" else d["location"]) or "未分類"
            groups.setdefault(str(key), []).append(d)
        group_count = len(groups)

        for key, group_defects in groups.items():
            sheet_name = unique_sheet_name(f"G-{key}", used_names)
            ws = wb.create_sheet(sheet_name)
            setup_photo_sheet_columns(ws)
            r = 1
            for d in group_defects:
                photos = defect_photo_map[d["id"]]
                anchor_row = r
                r = write_defect_block(ws, r, d, photos) + 2
                link_map[d["id"]] = (sheet_name, f"A{anchor_row}", len(photos))

    update_main_sheet_links(main_ws, defects, link_map)

    out_path = resolve_output_path(str(excel_path), args.output)
    try:
        wb.save(out_path)
    except PermissionError:
        die(f"無法寫入輸出檔案（可能已在 Excel 中開啟）：{out_path}")

    print_summary(
        excel_path, photos_root, out_path, defects, defect_photo_map, photo_count,
        total_orig, total_new, sheet_mode, group_count, warnings, errors,
    )


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as e:  # 任何未預期的錯誤都要給出明確訊息，不能直接崩潰
        if DEBUG:
            traceback.print_exc()
        else:
            print(f"\n❌ 發生未預期的錯誤：{e}\n（加上 --debug 參數可看到完整錯誤堆疊）", file=sys.stderr)
        sys.exit(1)
